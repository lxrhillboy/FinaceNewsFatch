#!/usr/bin/env python3
"""抓取习近平访美相关报道，保存到仓库根目录 output/日期/。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import List, Optional
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt

from news_fetcher.http import HttpClient
from news_fetcher.utils import clean_html_text

REPO_ROOT = Path(__file__).resolve().parents[1]


def resolve_output_dir(output_date: date) -> Path:
    return REPO_ROOT / "output" / output_date.isoformat()

# 访美行程核心报道（含 9 月 23 日当日及官方行程说明）
ARTICLE_URLS: List[dict] = [
    {
        "url": "https://www.news.cn/politics/leaders/20260923/1fe699a24ba84e819958dbb16812f358/c.html",
        "source": "新华网",
        "note": "2026-09-23 近镜头报道",
    },
    {
        "url": "https://www.gov.cn/yaowen/liebiao/202609/content_7081860.htm",
        "source": "中国政府网（新华社）",
        "note": "2026-09-23 习近平主席引领推进中美人民友好事业",
    },
    {
        "url": "https://www.wenweipo.com/a/202609/23/AP6ab32d04e4b01d54a2846756.html",
        "source": "香港文汇报",
        "note": "2026-09-23 特朗普接机欢迎习近平",
    },
    {
        "url": "https://www.gov.cn/yaowen/liebiao/202609/content_7081763.htm",
        "source": "中国政府网（新华社）",
        "note": "2026-09-22 美国各界对国事访问充满期待",
    },
    {
        "url": "https://www.gov.cn/yaowen/liebiao/202609/content_7081721.htm",
        "source": "中国政府网（新华社）",
        "note": "2026-09-22 共同擘画中美关系新篇章",
    },
    {
        "url": "https://www.gov.cn/lianbo/202609/content_7081659.htm",
        "source": "中国政府网（新华社）",
        "note": "2026-09-21 外交部介绍访美安排",
    },
    {
        "url": "https://www.news.cn/world/20260921/ad993f7f0e8c4404bb29c96d7a99e567/c.html",
        "source": "新华网",
        "note": "2026-09-21 外交部介绍访美安排",
    },
    {
        "url": "https://www.news.cn/20260921/cd3fe6083fd04971b767e11b49e2a3d6/c.html",
        "source": "新华网",
        "note": "2026-09-21 习近平将对美国进行国事访问",
    },
    {
        "url": "https://www.news.cn/politics/leaders/20260922/1b2f57b63dda444ea9a3d9b1491eafb0/c.html",
        "source": "新华网",
        "note": "2026-09-22 美国各界期待特稿",
    },
    {
        "url": "https://english.news.cn/20260922/a0b767faadce4259b47780b71da60874/c.html",
        "source": "新华社英文网",
        "note": "2026-09-22 Xinhua Headlines",
    },
    {
        "url": "https://english.www.gov.cn/news/202609/21/content_WS6ab0d84bc6d00ca5f9a0d4b0.html",
        "source": "中国政府网英文",
        "note": "2026-09-21 访美公告",
    },
    {
        "url": "https://www.whitehouse.gov/briefings-statements/2026/09/first-lady-melania-trump-releases-details-ahead-of-his-excellency-xi-jinping-president-of-the-peoples-republic-of-china-and-madame-peng-liyuans-visit-to-the-white-house/",
        "source": "白宫官网",
        "note": "2026-09-21 国事访问接待安排",
    },
    {
        "url": "https://www.voachinese.com/a/white-house-released-details-for-xi-visit-20260921/8202378.html",
        "source": "美国之音中文网",
        "note": "2026-09-21 白宫迎接细节",
    },
    {
        "url": "https://www.zaobao.com.sg/news/china/story20260922-9713681",
        "source": "联合早报",
        "note": "2026-09-22 白宫公布访美安排",
    },
    {
        "url": "https://www.zaobao.com.sg/news/china/story20260921-9709941",
        "source": "联合早报",
        "note": "2026-09-21 习近平9月23日至25日访美",
    },
    {
        "url": "https://www.bbc.com/zhongwen/articles/cx980mjemv73o/simp",
        "source": "BBC News 中文",
        "note": "2026-09-19 礼宾细节与行程",
    },
    {
        "url": "https://www.reuters.com/world/china/chinas-xi-visit-us-september-23-25-2026-09-21/",
        "source": "路透社",
        "note": "2026-09-21 访美日期报道",
        "fallback_title": "China's Xi to visit the US from September 23-25",
        "fallback_published": "2026-09-21",
        "fallback_body": (
            "BEIJING, Sept 21 (Reuters) — Chinese President Xi Jinping will pay a state "
            "visit to the US from September 23 to 25 at the invitation of President "
            "Donald Trump, China's Foreign Ministry said in a statement on Monday.\n\n"
            "During the visit, Xi \"will hold in-depth exchanges of views with President "
            "Trump on major issues concerning China-US relations, as well as world peace "
            "and development,\" a Chinese foreign ministry spokesperson said during a "
            "regular briefing.\n\n"
            "（注：路透社页面禁止自动化抓取，正文据公开电讯稿摘要录入，请以原文为准。）"
        ),
    },
    {
        "url": "https://bnonews.com/whpool/IRBL2_ZW",
        "source": "White House Press Pool (BNO News)",
        "note": "2026-09-21 白宫第一夫人办公室访美接待说明（转载）",
        "fallback_title": "First Lady Melania Trump Releases Details Ahead of Xi Jinping Visit",
        "fallback_published": "2026-09-21",
        "fallback_body": (
            "President Donald J. Trump and First Lady Melania Trump will welcome His "
            "Excellency Xi Jinping, President of the People's Republic of China and "
            "Madame Peng Liyuan to the White House on Thursday, September 24, 2026, "
            "for an official state visit.\n\n"
            "On Wednesday, September 23, the President will greet President Xi and "
            "Madame Peng at Joint Base Andrews.\n\n"
            "On Thursday morning, the President and First Lady will host an official "
            "arrival ceremony on the South Lawn, State Floor, and Rose Garden. On "
            "Thursday evening, the President and First Lady will greet President Xi "
            "and Madame Peng at the North Portico for a State Dinner in the East Room.\n\n"
            "On Friday morning, the President and First Lady will host a tea in the "
            "Red Room. The leaders will then visit the National Archives.\n\n"
            "（注：BNO 镜像页可能失效，正文与白宫官网声明一致，详见同目录白宫原文归档。）"
        ),
    },
]


@dataclass
class ImageAsset:
    original_url: str
    local_path: str
    alt: str = ""


@dataclass
class ArchivedArticle:
    title: str
    url: str
    source: str
    note: str
    published_at: str
    body_text: str
    body_html_excerpt: str
    images: List[ImageAsset] = field(default_factory=list)
    fetch_status: str = "ok"
    error: str = ""


def slugify(text: str, max_len: int = 60) -> str:
    text = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", text.strip())
    text = re.sub(r"-+", "-", text).strip("-")
    return (text or "article")[:max_len]


def extract_article_node(soup: BeautifulSoup):
    selectors = [
        "#UCAP-CONTENT",
        "div.pages_content",
        "div#detail",
        "div.detail",
        "article",
        "div.article-content",
        "div.story-content",
        "div.main-content",
        "div[data-component='text-block']",
        "div.entry-content",
        "div#content",
    ]
    for sel in selectors:
        node = soup.select_one(sel)
        if node and len(node.get_text(strip=True)) > 80:
            return node
    for tag in ("article", "main"):
        node = soup.find(tag)
        if node and len(node.get_text(strip=True)) > 120:
            return node
    return soup.body or soup


def parse_published(soup: BeautifulSoup, html: str) -> str:
    for sel in (
        "meta[property='article:published_time']",
        "meta[name='publishdate']",
        "meta[name='publishDate']",
        "meta[name='date']",
    ):
        tag = soup.select_one(sel)
        if tag and tag.get("content"):
            return tag["content"].strip()
    time_tag = soup.find("time")
    if time_tag:
        return (time_tag.get("datetime") or time_tag.get_text(strip=True) or "").strip()
    m = re.search(r"(\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2})", html)
    if m:
        return m.group(1)
    m = re.search(r"(\d{4}-\d{2}-\d{2})", html)
    if m:
        return m.group(1)
    return ""


def collect_images(node, page_url: str) -> List[tuple[str, str]]:
    images: List[tuple[str, str]] = []
    for img in node.find_all("img"):
        src = (img.get("src") or img.get("data-src") or img.get("data-original") or "").strip()
        if not src or src.startswith("data:"):
            continue
        full = urljoin(page_url, src)
        alt = (img.get("alt") or "").strip()
        images.append((full, alt))
    return images


def download_image(client: HttpClient, img_url: str, dest: Path) -> bool:
    try:
        resp = client.get(img_url)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(resp.content)
        return True
    except Exception:
        return False


def image_extension(url: str, content_type: str = "") -> str:
    path = urlparse(url).path
    ext = Path(path).suffix.lower()
    if ext in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg"}:
        return ext
    if "jpeg" in content_type:
        return ".jpg"
    if "png" in content_type:
        return ".png"
    if "webp" in content_type:
        return ".webp"
    return ".jpg"


def archive_one(
    client: HttpClient,
    meta: dict,
    articles_dir: Path,
    images_dir: Path,
    output_dir: Path,
) -> ArchivedArticle:
    url = meta["url"]
    source = meta.get("source", "")
    note = meta.get("note", "")
    try:
        html = client.get_text(url)
        soup = BeautifulSoup(html, "lxml")
        title_tag = soup.find("h1") or soup.find("title")
        title = title_tag.get_text(strip=True) if title_tag else url
        published = parse_published(soup, html)
        node = extract_article_node(soup)
        body_text = clean_html_text(node.decode_contents())
        body_html_excerpt = node.decode_contents()[:8000]

        article_slug = slugify(title)
        article_images_dir = images_dir / article_slug
        images: List[ImageAsset] = []
        seen_urls = set()
        for idx, (img_url, alt) in enumerate(collect_images(node, url), start=1):
            if img_url in seen_urls:
                continue
            seen_urls.add(img_url)
            digest = hashlib.sha1(img_url.encode()).hexdigest()[:10]
            ext = image_extension(img_url)
            local_name = f"{idx:02d}_{digest}{ext}"
            local_path = article_images_dir / local_name
            rel_path = str(local_path.relative_to(output_dir))
            if download_image(client, img_url, local_path):
                images.append(ImageAsset(original_url=img_url, local_path=rel_path, alt=alt))
            else:
                images.append(ImageAsset(original_url=img_url, local_path="", alt=alt))

        return ArchivedArticle(
            title=title,
            url=url,
            source=source,
            note=note,
            published_at=published,
            body_text=body_text,
            body_html_excerpt=body_html_excerpt,
            images=images,
        )
    except Exception as exc:  # noqa: BLE001
        fb_title = meta.get("fallback_title", url)
        fb_body = meta.get("fallback_body", "")
        fb_pub = meta.get("fallback_published", "")
        if fb_body:
            return ArchivedArticle(
                title=fb_title,
                url=url,
                source=source,
                note=note,
                published_at=fb_pub,
                body_text=fb_body,
                body_html_excerpt="",
                fetch_status="fallback",
                error=str(exc),
            )
        return ArchivedArticle(
            title=url,
            url=url,
            source=source,
            note=note,
            published_at="",
            body_text="",
            body_html_excerpt="",
            fetch_status="error",
            error=str(exc),
        )


def write_article_markdown(article: ArchivedArticle, path: Path) -> None:
    lines = [
        f"# {article.title}",
        "",
        f"- **信息来源**：{article.source}",
        f"- **原文网址**：{article.url}",
        f"- **归档说明**：{article.note}",
        f"- **发布时间**：{article.published_at or '（页面未标注）'}",
        f"- **抓取状态**：{article.fetch_status}",
    ]
    if article.error:
        lines.append(f"- **错误信息**：{article.error}")
    lines.extend(["", "## 正文（完整文字）", "", article.body_text or "（未能提取正文）", "", "## 图片"])
    if not article.images:
        lines.append("")
        lines.append("（本篇报道页面未包含可下载的正文配图，或图片下载失败。）")
    else:
        for img in article.images:
            lines.append("")
            lines.append(f"### {img.alt or '配图'}")
            lines.append(f"- 原图地址：{img.original_url}")
            if img.local_path:
                lines.append(f"![{img.alt or '配图'}]({img.local_path})")
            else:
                lines.append("（图片未能下载到本地）")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_word_report(articles: List[ArchivedArticle], path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    doc.add_heading("习近平主席访美报道汇编", level=0)
    doc.add_paragraph(
        f"共 {len(articles)} 篇；生成时间 {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    for art in sorted(articles, key=lambda a: (a.published_at, a.title)):
        doc.add_heading(art.title, level=1)
        doc.add_paragraph(f"信息来源：{art.source}")
        doc.add_paragraph(f"原文网址：{art.url}")
        doc.add_paragraph(f"发布时间：{art.published_at or '（未标注）'}")
        doc.add_paragraph(f"抓取状态：{art.fetch_status}")
        doc.add_paragraph(art.body_text or "（未能提取正文）")
        if art.images:
            doc.add_paragraph("配图原址：")
            for img in art.images:
                doc.add_paragraph(img.original_url, style="List Bullet")
        doc.add_page_break()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def write_index(articles: List[ArchivedArticle], path: Path, output_date: date) -> None:
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    lines = [
        f"# 习近平主席访美报道归档（{output_date.isoformat()}）",
        "",
        f"本目录由仓库脚本自动抓取整理，输出路径：`output/{output_date.isoformat()}/`，"
        f"最后更新：**{now}**。",
        "",
        "## 访问背景",
        "",
        "应美国总统特朗普邀请，国家主席习近平于 **2026年9月23日至25日** 对美国进行国事访问。"
        "公开行程安排包括：9月23日在安德鲁斯联合基地抵达欢迎；9月24日白宫正式欢迎仪式、"
        "双边会谈与国宴；9月25日茶叙、参观美国国家档案馆后离境。",
        "",
        "## 报道索引",
        "",
        "| 发布时间 | 来源 | 标题 | 网址 | 本地文件 |",
        "| --- | --- | --- | --- | --- |",
    ]
    for art in sorted(articles, key=lambda a: (a.published_at, a.title)):
        slug = slugify(art.title)
        md_link = f"articles/{slug}.md"
        title = art.title.replace("|", "\\|")
        lines.append(
            f"| {art.published_at or '-'} | {art.source} | {title} | {art.url} | [{slug}.md]({md_link}) |"
        )
    lines.extend(
        [
            "",
            "## 数据文件",
            "",
            "- Word 汇总：`report.docx`",
            "- 结构化元数据：`metadata.json`",
            "- 单篇 Markdown：`articles/`",
            "- 配图：`images/`",
            "- 抓取脚本：`scripts/archive_xi_us_visit.py`",
            "",
            "## 说明",
            "",
            "部分境外媒体（如路透社）可能对自动化访问有限制；若 `fetch_status` 为 error，"
            "请直接打开原文网址阅读。图片版权归原媒体所有，本地副本仅用于研究与归档。",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="抓取习近平访美报道到 output/日期/")
    parser.add_argument(
        "--date",
        help="输出子目录日期，格式 YYYY-MM-DD（默认：当天 UTC 日期）",
        default=None,
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.date:
        output_date = datetime.strptime(args.date, "%Y-%m-%d").date()
    else:
        output_date = datetime.now(timezone.utc).date()

    output_dir = resolve_output_dir(output_date)
    client = HttpClient(timeout=35, retries=3)
    articles_dir = output_dir / "articles"
    images_dir = output_dir / "images"
    articles_dir.mkdir(parents=True, exist_ok=True)
    images_dir.mkdir(parents=True, exist_ok=True)

    archived: List[ArchivedArticle] = []
    for meta in ARTICLE_URLS:
        print(f"抓取: {meta['url']}")
        article = archive_one(client, meta, articles_dir, images_dir, output_dir)
        archived.append(article)
        md_path = articles_dir / f"{slugify(article.title)}.md"
        write_article_markdown(article, md_path)

    payload = {
        "topic": f"习近平主席访美 {output_date.isoformat()}",
        "output_dir": str(output_dir.relative_to(REPO_ROOT)),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "article_count": len(archived),
        "articles": [
            {
                **{k: v for k, v in asdict(a).items() if k != "images"},
                "images": [asdict(img) for img in a.images],
            }
            for a in archived
        ],
    }
    (output_dir / "metadata.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_index(archived, output_dir / "README.md", output_date)
    write_word_report(archived, output_dir / "report.docx")
    ok = sum(1 for a in archived if a.fetch_status == "ok")
    print(f"输出目录: {output_dir}")
    print(f"完成：{ok}/{len(archived)} 篇成功")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
